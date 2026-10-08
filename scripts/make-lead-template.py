"""
Generates public/templates/salespulse-lead-template.xlsx (+ .csv).
Run: python3 scripts/make-lead-template.py   (needs openpyxl)
Keep the column names in sync with lib/leadColumns.ts.
"""
import csv, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = os.path.join(os.path.dirname(__file__), "..", "public", "templates")
os.makedirs(OUT, exist_ok=True)

COLUMNS = [
    ("Lead Name *", 26, "Required. The person you'll be calling, e.g. John Banda."),
    ("Phone Number *", 18, "Required. Local (0977 123456) or international (+260 977 123456). If there are two numbers, separate them with  /  — the first is used, the other is kept in notes."),
    ("Company", 26, "Optional. Business or organisation, if any."),
    ("Email", 28, "Optional."),
    ("Industry", 18, "Optional, e.g. Logistics, Healthcare, Retail."),
    ("Location", 18, "Optional. Town / city / area."),
    ("Website", 22, "Optional."),
    ("Product", 20, "Optional. Product of interest — use the same name as in SalesPulse (e.g. Visit Flow, Marichi Pulse, Guardly)."),
    ("Source", 20, "Optional. Pick from the list: Google Search, VisitFlow, Meta Leads, Referral, Website, Other."),
    ("Notes", 36, "Optional. Anything the caller should know."),
]
SOURCES = ["Google Search", "VisitFlow", "Meta Leads", "Referral", "Website", "Other"]

PURPLE = "5B21B6"
LIGHT = "EDE9FE"
REQ = "FDE68A"
thin = Side(style="thin", color="D9D4EC")

wb = Workbook()
ws = wb.active
ws.title = "Leads"
for i, (name, width, _) in enumerate(COLUMNS, start=1):
    c = ws.cell(row=1, column=i, value=name)
    required = name.endswith("*")
    c.font = Font(bold=True, color="FFFFFF" if not required else "1F1300")
    c.fill = PatternFill("solid", fgColor=PURPLE if not required else REQ)
    c.alignment = Alignment(vertical="center")
    c.border = Border(bottom=thin)
    ws.column_dimensions[c.column_letter].width = width
ws.row_dimensions[1].height = 22
ws.freeze_panes = "A2"
# Phone column as text so Excel keeps leading zeros (0977…), for the first 5,000 rows.
for r in range(2, 5002):
    ws.cell(row=r, column=2).number_format = "@"
dv = DataValidation(type="list", formula1='"' + ",".join(SOURCES) + '"', allow_blank=True, showErrorMessage=False)
ws.add_data_validation(dv)
dv.add("I2:I5001")

ins = wb.create_sheet("Instructions")
ins.column_dimensions["A"].width = 22
ins.column_dimensions["B"].width = 95
ins["A1"] = "SalesPulse — lead upload template"
ins["A1"].font = Font(bold=True, size=14, color=PURPLE)
rows = [
    ("How to use", "Fill in the 'Leads' sheet — one lead per row — and upload it in SalesPulse: Leads → Bulk upload (or My Day → Bulk upload)."),
    ("Required", "Only two columns are required: Lead Name and Phone Number (yellow headers). Everything else is optional — leave it blank if you don't have it."),
    ("Duplicates", "Numbers already in SalesPulse, or repeated in the same file, are skipped automatically — you'll see a preview before anything is saved."),
    ("Other formats", "You don't have to use this exact layout: any .xlsx, .xls or .csv with a name column and a phone column works — SalesPulse lets you match the columns when you upload."),
    ("Limit", "Up to 5,000 rows per file."),
]
for i, (k, v) in enumerate(rows, start=3):
    ins.cell(row=i, column=1, value=k).font = Font(bold=True)
    ins.cell(row=i, column=2, value=v).alignment = Alignment(wrap_text=True, vertical="top")
r0 = 3 + len(rows) + 1
ins.cell(row=r0, column=1, value="Columns").font = Font(bold=True, size=12, color=PURPLE)
for i, (name, _, desc) in enumerate(COLUMNS, start=r0 + 1):
    ins.cell(row=i, column=1, value=name).font = Font(bold=name.endswith("*"))
    ins.cell(row=i, column=2, value=desc).alignment = Alignment(wrap_text=True, vertical="top")
r1 = r0 + len(COLUMNS) + 2
ins.cell(row=r1, column=1, value="Example rows").font = Font(bold=True, size=12, color=PURPLE)
example = [
    ["John Banda", "0977 123456", "ABC Logistics Ltd", "john@abclogistics.co.zm", "Logistics", "Lusaka", "abclogistics.co.zm", "Visit Flow", "Google Search", "Asked for a call after 2pm"],
    ["Mary Phiri", "+260 966 111222", "", "", "", "Ndola", "", "", "Referral", ""],
    ["Peter Mwale", "0955 333444 / 0977 555666", "", "", "", "", "", "", "Meta Leads", ""],
]
hdr = ins.cell(row=r1 + 1, column=1)
for j, (name, _, _) in enumerate(COLUMNS, start=1):
    c = ins.cell(row=r1 + 1, column=j, value=name)
    c.font = Font(bold=True)
    c.fill = PatternFill("solid", fgColor=LIGHT)
for i, row in enumerate(example, start=r1 + 2):
    for j, v in enumerate(row, start=1):
        ins.cell(row=i, column=j, value=v)
for col in "CDEFGHIJ":
    ins.column_dimensions[col].width = 18

wb.save(os.path.join(OUT, "salespulse-lead-template.xlsx"))
with open(os.path.join(OUT, "salespulse-lead-template.csv"), "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    w.writerow([c[0] for c in COLUMNS])
print("written", OUT)
